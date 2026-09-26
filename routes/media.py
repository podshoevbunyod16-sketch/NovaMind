"""
routes/media.py — Загрузка файлов, изображений, TTS, STT, видео
"""
from flask import Blueprint, request, jsonify, send_from_directory, send_file, session
import os
import uuid
import mimetypes
import base64
import zipfile
import xml.etree.ElementTree as ET
import requests
from ai_providers import groq_request_with_rotation
from groq_rotation import get_groq_key
import config

media_bp = Blueprint("media", __name__)

# ========== ОТДАЧА ИЗОБРАЖЕНИЙ ==========

@media_bp.route('/generated_image')
def generated_image():
    filename = request.args.get("file", "")
    if not filename:
        return jsonify({"error": "No filename"}), 400
    safe_name = filename.replace("..", "").replace("/", "")
    img_dir = os.path.join(os.path.dirname(__file__), "generated_images")
    filepath = os.path.join(img_dir, safe_name)
    if not os.path.exists(filepath):
        return jsonify({"error": "File not found"}), 404
    return send_file(filepath, mimetype='image/png')


# ========== ЗАГРУЗКА ФАЙЛОВ ==========


@media_bp.route('/upload_image', methods=['POST'])
def upload_image():
    return analyze_attachment()

@media_bp.route('/upload_file', methods=['POST'])
def upload_file():
    return analyze_attachment()

@media_bp.route('/api/attachments/analyze', methods=['POST'])
def analyze_attachment():
    """Universal analyzer for images and common documents."""
    global contents
    uploaded = request.files.get('file') or request.files.get('image')
    if not uploaded or not uploaded.filename:
        return jsonify({'error': 'Файл не выбран'}), 400

    filename = os.path.basename(uploaded.filename)
    ext = os.path.splitext(filename)[1].lower()
    content_type = uploaded.mimetype or mimetypes.guess_type(filename)[0] or 'application/octet-stream'
    user_desc = (request.form.get('description') or '').strip()

    upload_dir = os.path.join(os.path.dirname(__file__), "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    filepath = os.path.join(upload_dir, f"{uuid.uuid4().hex}_{filename}")
    uploaded.save(filepath)

    try:
        if os.path.getsize(filepath) > 25 * 1024 * 1024:
            return jsonify({'error': 'Файл слишком большой. Максимальный размер анализа — 25 МБ.'}), 413

        is_image = content_type.startswith('image/') or ext in {'.jpg','.jpeg','.png','.webp','.gif','.bmp'}
        if is_image:
            return _analyze_image_file(filepath, filename, content_type, user_desc)

        text = _extract_attachment_text(filepath, ext, content_type)
        if not text.strip():
            return jsonify({'error': f'Файл {filename} пустой или из него не удалось извлечь текст.'}), 415

        instruction = user_desc or _default_file_instruction(ext)
        prompt = (
            f"Файл: {filename}\n\nЗадача: {instruction}\n\n"
            f"Содержимое:\n{text[:120000]}\n\n"
            "Отвечай на русском языке в Markdown. Не выдумывай данные. "
            "Для кода указывай конкретные ошибки и исправления."
        )
        reply, error = _call_selected_text_ai(prompt)
        if error:
            return jsonify({'error': error}), 502

        return jsonify({
            'success': True,
            'filename': filename,
            'type': content_type,
            'result': f'📁 **Анализ файла {filename}:**\n\n{reply}'
        })
    except Exception as exc:
        print(f"[attachments] error: {exc}")
        return jsonify({'error': f'Ошибка обработки файла: {exc}'}), 500


def _default_file_instruction(ext):
    if ext == '.py':
        return 'Проанализируй Python-код, объясни его работу, найди ошибки и предложи улучшения.'
    if ext in {'.js','.ts','.jsx','.tsx'}:
        return 'Проанализируй JavaScript/TypeScript-код, найди ошибки и предложи улучшения.'
    if ext in {'.html','.css'}:
        return 'Проанализируй веб-файл, объясни структуру и найди проблемы.'
    if ext == '.json':
        return 'Проанализируй структуру JSON и объясни важные данные.'
    if ext in {'.csv','.tsv'}:
        return 'Проанализируй таблицу, колонки, данные и важные закономерности.'
    if ext == '.md':
        return 'Сделай структурированное резюме документа и выдели главное.'
    return 'Подробно проанализируй файл и объясни его содержимое.'


def _extract_attachment_text(filepath, ext, content_type):
    text_exts = {
        '.txt','.json','.csv','.tsv','.py','.js','.ts','.jsx','.tsx','.html','.htm',
        '.css','.md','.xml','.yaml','.yml','.log','.ini','.cfg','.conf','.sql',
        '.sh','.bash','.java','.c','.cpp','.h','.hpp','.go','.rs','.php'
    }
    if ext in text_exts or content_type.startswith('text/'):
        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
            return f.read()

    if ext == '.pdf':
        try:
            from pypdf import PdfReader
            return '\n\n'.join((page.extract_text() or '') for page in PdfReader(filepath).pages)
        except ImportError:
            raise RuntimeError('PDF: зависимость pypdf не установлена.')

    if ext == '.docx':
        try:
            from docx import Document
            doc = Document(filepath)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(' | '.join(cell.text.strip() for cell in row.cells))
            return '\n'.join(parts)
        except ImportError:
            raise RuntimeError('DOCX: зависимость python-docx не установлена.')

    if ext in {'.xlsx','.xlsm'}:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(filepath, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f'### Лист: {ws.title}')
                for row in ws.iter_rows(values_only=True):
                    values = [str(v) if v is not None else '' for v in row]
                    if any(values):
                        parts.append(' | '.join(values))
            return '\n'.join(parts)
        except ImportError:
            raise RuntimeError('XLSX: зависимость openpyxl не установлена.')

    if ext == '.pptx':
        with zipfile.ZipFile(filepath) as z:
            parts = []
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for name in sorted(z.namelist()):
                if name.startswith('ppt/slides/slide') and name.endswith('.xml'):
                    root = ET.fromstring(z.read(name))
                    parts.extend(t.text for t in root.findall('.//a:t', ns) if t.text)
            return '\n'.join(parts)

    raise RuntimeError(
        f'Формат {ext or content_type} не поддерживается. '
        'Поддерживаются изображения, TXT/код/JSON/CSV, PDF, DOCX, XLSX/XLSM и PPTX.'
    )


def _analyze_image_file(filepath, filename, content_type, user_desc):
    instruction = user_desc or 'Подробно опиши изображение: объекты, текст, структуру и важные детали.'
    with open(filepath, 'rb') as f:
        encoded = base64.b64encode(f.read()).decode('ascii')

    # Gemini Vision: основной путь.
    gemini_key = (os.getenv('GEMINI_API_KEY') or os.getenv('GOOGLE_AI_STUDIO_KEY') or '').strip()
    if gemini_key:
        model = os.getenv('VISION_MODEL', 'gemini-2.5-flash')
        body = {
            'contents': [{'role': 'user', 'parts': [
                {'text': instruction},
                {'inline_data': {'mime_type': content_type, 'data': encoded}}
            ]}],
            'generationConfig': {'temperature': 0.2, 'maxOutputTokens': 4096}
        }
        try:
            r = requests.post(
                f'https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent',
                params={'key': gemini_key}, json=body,
                headers={'Content-Type': 'application/json'}, timeout=120
            )
            if r.ok:
                parts = []
                for candidate in r.json().get('candidates', []):
                    for part in (candidate.get('content') or {}).get('parts', []):
                        if part.get('text'):
                            parts.append(part['text'])
                if parts:
                    return jsonify({'success': True, 'filename': filename,
                        'result': f'📷 **Анализ изображения {filename} (Google AI Studio):**\n\n{"".join(parts)}'})
            print(f"[attachments] Gemini Vision HTTP {r.status_code}: {r.text[:500]}")
        except requests.RequestException as exc:
            print(f"[attachments] Gemini Vision error: {exc}")

    # Groq Vision fallback.
    groq_key = get_groq_key()
    if groq_key:
        payload = {
            'model': 'meta-llama/llama-4-scout-17b-16e-instruct',
            'messages': [{'role':'user','content':[
                {'type':'text','text':instruction},
                {'type':'image_url','image_url':{'url':f'data:{content_type};base64,{encoded}'}}
            ]}],
            'max_tokens': 4096
        }
        try:
            r = requests.post('https://api.groq.com/openai/v1/chat/completions',
                headers={'Authorization':f'Bearer {groq_key}','Content-Type':'application/json'},
                json=payload, timeout=120)
            r.raise_for_status()
            reply = r.json()['choices'][0]['message']['content']
            return jsonify({'success': True, 'filename': filename,
                'result': f'📷 **Анализ изображения {filename} (Vision):**\n\n{reply}'})
        except Exception as exc:
            print(f"[attachments] Groq Vision error: {exc}")

    return jsonify({'error': 'Нет доступного Vision-провайдера. Добавьте GEMINI_API_KEY/GOOGLE_AI_STUDIO_KEY или GROQ API key.'}), 503


def _call_selected_text_ai(prompt):
    provider = config.PROVIDERS.get(config.current_provider)
    if not provider:
        return None, f'Неизвестный AI-провайдер: {config.current_provider}'
    payload = {
        'model': config.current_model,
        'messages': [
            {'role': 'system', 'content': config.system_prompt},
            {'role': 'user', 'content': prompt}
        ],
        'temperature': 0.3,
        'max_tokens': min(provider.get('max_tokens', 8192), 8192)
    }
    data, error = groq_request_with_rotation(
        provider.get('url',''), payload, provider.get('headers',{}).copy(), timeout=120
    )
    if error:
        return None, error
    try:
        return data['choices'][0]['message']['content'], None
    except (KeyError, IndexError, TypeError):
        return None, 'AI вернул неожиданный формат ответа.'


# ========== МУЛЬТИМЕДИА ==========
MEDIA_DIR = os.path.join(os.path.dirname(__file__), "generated_media")
os.makedirs(MEDIA_DIR, exist_ok=True)

@media_bp.route('/media/<path:filename>')
def media_file(filename):
    safe_name = os.path.basename(filename)
    filepath = os.path.join(MEDIA_DIR, safe_name)
    if not os.path.isfile(filepath):
        return jsonify({"error": "Медиафайл не найден"}), 404
    return send_file(filepath, mimetype=mimetypes.guess_type(filepath)[0] or "application/octet-stream")

@media_bp.route('/api/media/upload', methods=['POST'])
def media_upload():
    media = request.files.get('file')
    if not media or not media.filename:
        return jsonify({"error": "Файл не выбран"}), 400
    content_type = media.mimetype or mimetypes.guess_type(media.filename)[0] or "application/octet-stream"
    if not (content_type.startswith("audio/") or content_type.startswith("video/") or content_type.startswith("image/")):
        return jsonify({"error": "Поддерживаются только аудио, видео и изображения"}), 415
    extension = os.path.splitext(media.filename)[1].lower() or mimetypes.guess_extension(content_type) or ""
    saved_name = f"upload_{uuid.uuid4().hex}{extension}"
    media.save(os.path.join(MEDIA_DIR, saved_name))
    return jsonify({"success": True, "filename": media.filename, "type": content_type, "url": f"/media/{saved_name}"})

@media_bp.route('/api/media/transcribe', methods=['POST'])
def media_transcribe():
    media = request.files.get('file')
    if not media or not media.filename:
        return jsonify({"error": "Аудиофайл не выбран"}), 400
    endpoint = os.getenv("AUDIO_TRANSCRIPTION_URL", "")
    api_key = os.getenv("AUDIO_API_KEY", "")
    if not endpoint:
        if os.getenv("GROQ_API_KEY") or GROQ_KEYS:
            endpoint = "https://api.groq.com/openai/v1/audio/transcriptions"
            api_key = api_key or get_groq_key()
        elif os.getenv("OPENAI_API_KEY"):
            endpoint = "https://api.openai.com/v1/audio/transcriptions"
            api_key = api_key or os.getenv("OPENAI_API_KEY", "")
    if not endpoint:
        return jsonify({"error": "Настройте AUDIO_TRANSCRIPTION_URL и AUDIO_API_KEY либо добавьте GROQ_API_KEY/OPENAI_API_KEY в .env"}), 503
    headers = {"Authorization": f"Bearer {api_key}"} if api_key else {}
    try:
        response = requests.post(endpoint, headers=headers, files={"file": (media.filename, media.stream, media.mimetype)}, data={"model": os.getenv("AUDIO_TRANSCRIPTION_MODEL", "whisper-large-v3-turbo"), "language": request.form.get("language", "ru")}, timeout=180)
        response.raise_for_status()
        result = response.json()
        return jsonify({"success": True, "text": result.get("text", ""), "language": result.get("language")})
    except (requests.RequestException, ValueError) as exc:
        return jsonify({"error": f"Ошибка расшифровки аудио: {exc}"}), 502

@media_bp.route('/api/media/tts', methods=['POST'])
def media_tts():
    data = request.get_json() or {}
    text = data.get("text", "").strip()
    if not text:
        return jsonify({"error": "Введите текст для озвучивания"}), 400
    endpoint = os.getenv("AUDIO_TTS_URL", "")
    api_key = os.getenv("AUDIO_API_KEY", "")
    model = os.getenv("AUDIO_TTS_MODEL", "tts-1")
    if not endpoint and os.getenv("OPENAI_API_KEY"):
        endpoint, api_key = "https://api.openai.com/v1/audio/speech", api_key or os.getenv("OPENAI_API_KEY", "")
    if not endpoint:
        return jsonify({"error": "Настройте AUDIO_TTS_URL и AUDIO_API_KEY либо добавьте OPENAI_API_KEY в .env"}), 503
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"} if api_key else {"Content-Type": "application/json"}
    try:
        response = requests.post(endpoint, headers=headers, json={"model": model, "input": text, "voice": data.get("voice", "alloy"), "response_format": "mp3"}, timeout=180)
        response.raise_for_status()
        filename = f"speech_{uuid.uuid4().hex}.mp3"
        with open(os.path.join(MEDIA_DIR, filename), "wb") as audio_file:
            audio_file.write(response.content)
        return jsonify({"success": True, "url": f"/media/{filename}", "filename": filename})
    except requests.RequestException as exc:
        return jsonify({"error": f"Ошибка генерации аудио: {exc}"}), 502

@media_bp.route('/api/media/image', methods=['POST'])
def media_image():
    data = request.get_json() or {}
    prompt = data.get("prompt", "").strip()
    if not prompt:
        return jsonify({"error": "Введите описание изображения"}), 400
    endpoint = os.getenv("IMAGE_GENERATION_URL", "")
    api_key = os.getenv("IMAGE_API_KEY", "") or os.getenv("OPENAI_API_KEY", "")
    if not endpoint and os.getenv("OPENAI_API_KEY"):
        endpoint = "https://api.openai.com/v1/images/generations"
    if not endpoint:
        from urllib.parse import quote
        width, height = int(data.get("width", 1024)), int(data.get("height", 1024))
        url = f"https://image.pollinations.ai/prompt/{quote(prompt)}?width={width}&height={height}&nologo=True"
        return jsonify({"success": True, "url": url, "provider": "pollinations"})
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"} if api_key else {"Content-Type": "application/json"}
    try:
        response = requests.post(endpoint, headers=headers, json={"model": os.getenv("IMAGE_MODEL", "gpt-image-1"), "prompt": prompt, "size": data.get("size", "1024x1024"), "n": 1}, timeout=240)
        response.raise_for_status()
        result = response.json().get("data", [{}])[0]
        if result.get("url"):
            return jsonify({"success": True, "url": result["url"], "provider": "configured"})
        if result.get("b64_json"):
            import base64
            filename = f"image_{uuid.uuid4().hex}.png"
            with open(os.path.join(MEDIA_DIR, filename), "wb") as image_file:
                image_file.write(base64.b64decode(result["b64_json"]))
            return jsonify({"success": True, "url": f"/media/{filename}", "provider": "configured"})
        return jsonify({"error": "Провайдер не вернул изображение"}), 502
    except (requests.RequestException, ValueError, IndexError) as exc:
        return jsonify({"error": f"Ошибка генерации изображения: {exc}"}), 502

@media_bp.route('/api/media/video', methods=['POST'])
def media_video():
    data = request.get_json() or {}
    prompt = data.get("prompt", "").strip()
    endpoint = os.getenv("VIDEO_API_URL", "")
    if not prompt:
        return jsonify({"error": "Введите описание видео"}), 400
    if not endpoint:
        return jsonify({"error": "Для видео укажите VIDEO_API_URL, VIDEO_API_KEY и VIDEO_MODEL в .env. У видео-провайдеров разные API, поэтому endpoint настраивается явно."}), 503
    api_key = os.getenv("VIDEO_API_KEY", "") or os.getenv("OPENAI_API_KEY", "")
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"} if api_key else {"Content-Type": "application/json"}
    payload = {"model": os.getenv("VIDEO_MODEL", ""), "prompt": prompt, "duration": data.get("duration", 5), "aspect_ratio": data.get("aspect_ratio", "16:9")}
    try:
        response = requests.post(endpoint, headers=headers, json=payload, timeout=240)
        response.raise_for_status()
        result = response.json()
        return jsonify({"success": True, "url": result.get("video_url") or result.get("url"), "job_id": result.get("id") or result.get("job_id"), "status_url": result.get("status_url"), "raw": result})
    except (requests.RequestException, ValueError) as exc:
        return jsonify({"error": f"Ошибка генерации видео: {exc}"}), 502

# ========== МОДЕЛИ ==========
